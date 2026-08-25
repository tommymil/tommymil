"""
Zamienia hurtowy cennik XLSX dostawcy na dane katalogu sklepu.

Wynik:
  * shop-catalog.json  — pozycje gotowe do wczytania przez ApplicationDbContextSeeder,
  * katalog z obrazami — zdjęcia przekonwertowane do WebP, po jednym pliku na obraz.

Ceny w arkuszu są hurtowe netto. Cena sklepowa to netto * (1 + marża) * (1 + VAT),
zaokrąglona do pełnych groszy — sklep trzyma kwoty brutto w groszach jako int.

Arkusz nie ma kodów produktów, więc kluczem jest slug wyprowadzony z nazwy. Przy kolejnej
wersji cennika zmiana nazwy zerwie powiązanie — to ograniczenie źródła, nie skryptu.

Użycie:
    python import_cennik.py "C:/.../Cennik hurtowy.xlsx"
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import unicodedata
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from PIL import Image

MARGIN = Decimal("0.10")   # marża A.K. HOUSE na cenie hurtowej netto
VAT = Decimal("0.23")

SHEET = "POLISH"
COL_NAME = 5               # F — nazwa + opis
COL_PRICE = 6              # G — cena netto
COL_LINK = 9               # J — link do oferty dostawcy
SPEC_COLS = range(0, 5)    # A–E — zdjęcia i dane techniczne

MAX_NAME = 200             # limity z ShopProductConfiguration
MAX_SHORT = 400
MAX_DESCRIPTION = 8000
MAX_SLUG = 160
MAX_IMAGES = 4
IMAGE_MAX_WIDTH = 1400
IMAGE_QUALITY = 82

# Sekcje cennika -> półki sklepu. Wartości to ShopCategory.Key z bazy (tabela ShopCategories),
# nie nazwy enuma — enum zniknął, gdy kategorie stały się edytowalne w /sklep/panel. Klucze
# celowo zostały te same, więc ten słownik nie wymagał zmiany, ALE: kategoria dodana w panelu
# ma klucz wyprowadzony ze slugu (np. "reczniki-i-tekstylia"). Chcesz importować do niej z cennika
# — wpisz tu dokładnie ten klucz, widoczny w panelu w wierszu kategorii.
# Nagłówki sekcji przepisane z arkusza;
# porównanie idzie po tekście znormalizowanym (bez ogonków, bez wielokrotnych spacji).
SECTION_CATEGORIES = {
    "akcesoria do wanien spa / basenowe": "Accessories",
    "termometry": "Accessories",
    "pozostale": "Accessories",
    "akryl standard ( szara perla, biala perla, bezowa perla)": "TubShells",
    "akryl standard bez wytloczen ( szara perla, biala perla, zlota perla)": "TubShells",
    "akryl z wytloczeniami ( szary metalik, bezowy metalik, bialy metalik)": "TubShells",
    "wklady z wlokna szklanego": "TubShells",
    "wklady z wlokna szklanego do beczek do schladzania": "TubShells",
    "piece i akcesoria": "Heaters",
    "elementy do pieca": "Heaters",
    "pokrywy termiczne": "Covers",
    "systemy masazu": "MassageSystems",
    "dysze": "MassageSystems",
    "oswietlenie": "Lighting",
    "naglosnienie": "Lighting",
    "elementy drewniane": "WoodenElements",
    "minibarki": "Minibars",
    "filtracja": "Filtration",
    "grzalki": "WaterHeaters",
    "chemia": "Chemicals",
    "elementy do montazu": "Assembly",
    # Obie sekcje „Węże" i druga „Odpływy" to w rzeczywistości złączki, obejmy i zawory PVC —
    # tytuły w cenniku się powtarzają i nie opisują zawartości. Wszystko idzie na jedną półkę.
    "weze": "Assembly",
    "odplywy": "Assembly",
}

# Pojedyncze wiersze, których sekcja nie opisuje. Klucz to numer wiersza w arkuszu.
ROW_CATEGORIES = {
    163: "Controls",  # sterownik do balii
    164: "Controls",  # sterownik Wi-Fi
    165: "Controls",  # zabezpieczenie różnicowo-prądowe
}

# Sekcje, w których sama nazwa pozycji („Okrągły 200 cm standard") nic nie mówi na liście.
SECTION_NAME_PREFIXES = {
    "wklady z wlokna szklanego": "Wkład z włókna szklanego",
    "wklady z wlokna szklanego do beczek do schladzania": "Wkład do beczki schładzającej",
    "akryl standard bez wytloczen ( szara perla, biala perla, zlota perla)": "Wkład akrylowy",
    "akryl z wytloczeniami ( szary metalik, bezowy metalik, bialy metalik)": "Wkład akrylowy z wytłoczeniami",
}

# Dopiski, którymi dostawca oznacza pozycje w arkuszu.
MARKER_NEW = re.compile(r"nowo[śs][ćc]\s*!+", re.IGNORECASE)
MARKER_UNAVAILABLE = re.compile(r"niedost[ęe]pn[ya]\s*!+", re.IGNORECASE)

# Wiersze z jedną komórką w kolumnie A, które są banerem, a nie nagłówkiem sekcji.
SECTION_BANNERS = {"polski producent !"}

# Jednostka rozliczeniowa -> dopisek do nazwy i zdania w opisie skróconym.
UNIT_LABELS = {
    "m": ("1 mb", "metr bieżący"),
    "mb": ("1 mb", "metr bieżący"),
    "m2": ("1 m²", "metr kwadratowy"),
    "szt": ("1 szt.", "sztukę"),
    "kpl": ("1 kpl.", "komplet"),
    "2szt": ("2 szt.", "parę"),
}

PL_CHARS = str.maketrans("ąćęłńóśźżĄĆĘŁŃÓŚŹŻ", "acelnoszzACELNOSZZ")

XDR = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
MAIN_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def normalise(text: str) -> str:
    """Tekst do porównań: bez ogonków, małe litery, pojedyncze spacje."""
    return re.sub(r"\s+", " ", text.translate(PL_CHARS).strip().lower())


def slugify(text: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", text.translate(PL_CHARS))
    ascii_text = ascii_text.encode("ascii", "ignore").decode("ascii").lower()
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_text).strip("-")
    return slug[:MAX_SLUG].rstrip("-")


def clean(value) -> str:
    """Komórka -> tekst bez podwójnych spacji i twardych spacji."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"[ \t\u00a0]+", " ", str(value)).strip()


def shorten(text: str, limit: int) -> str:
    """Skraca po granicy słowa; wielokropek tylko gdy faktycznie coś ucięto."""
    text = text.strip()
    if len(text) <= limit:
        return text
    cut = text[: limit - 1]
    space = cut.rfind(" ")
    if space > limit // 2:
        cut = cut[:space]
    return cut.rstrip(" ,.;:-") + "…"


@dataclass
class Product:
    slug: str
    name: str
    short_description: str
    description: str
    price_grosze: int
    wholesale_net_grosze: int
    category: str
    sort_order: int
    is_published: bool
    unit: str | None
    source_url: str | None
    source_row: int
    images: list[str] = field(default_factory=list)


def read_hyperlinks(archive: zipfile.ZipFile, sheet_index: int) -> dict[str, str]:
    """Adresy z kolumny J. Komórka pokazuje czasem sam tytuł oferty — URL siedzi w relacji."""
    rels_path = f"xl/worksheets/_rels/sheet{sheet_index}.xml.rels"
    targets = {
        rel.get("Id"): rel.get("Target")
        for rel in ET.fromstring(archive.read(rels_path))
        if rel.get("Type", "").endswith("/hyperlink")
    }

    links: dict[str, str] = {}
    sheet = ET.fromstring(archive.read(f"xl/worksheets/sheet{sheet_index}.xml"))
    for hyperlink in sheet.findall(".//m:hyperlink", MAIN_NS):
        target = targets.get(hyperlink.get(R_NS + "id"))
        if target:
            links[hyperlink.get("ref")] = target.strip()
    return links


def read_image_anchors(archive: zipfile.ZipFile, drawing_index: int) -> dict[int, list[tuple[int, str]]]:
    """Wiersz arkusza -> [(kolumna, ścieżka pliku w archiwum)], w kolejności kolumn."""
    rels_path = f"xl/drawings/_rels/drawing{drawing_index}.xml.rels"
    targets = {
        rel.get("Id"): rel.get("Target").replace("../", "xl/")
        for rel in ET.fromstring(archive.read(rels_path))
    }

    anchors: dict[int, list[tuple[int, str]]] = defaultdict(list)
    drawing = ET.fromstring(archive.read(f"xl/drawings/drawing{drawing_index}.xml"))
    for anchor in drawing:
        origin = anchor.find(XDR + "from")
        blip = anchor.find(f".//{{http://schemas.openxmlformats.org/drawingml/2006/main}}blip")
        if origin is None or blip is None:
            continue
        target = targets.get(blip.get(R_NS + "embed"))
        if not target:
            continue
        row = int(origin.find(XDR + "row").text) + 1
        column = int(origin.find(XDR + "col").text) + 1
        anchors[row].append((column, target))

    for row in anchors:
        anchors[row].sort(key=lambda item: item[0])
    return anchors


def parse_price(raw) -> tuple[Decimal, str | None]:
    """Cena netto i jednostka. Część wierszy ma cenę jako tekst: '64 zł / m2'."""
    if isinstance(raw, (int, float)):
        return Decimal(str(raw)), None

    text = clean(raw)
    match = re.match(r"([\d\s]+(?:[.,]\d+)?)\s*z[łl]\s*(?:/\s*(.+))?$", text, re.IGNORECASE)
    if not match:
        raise ValueError(f"nieczytelna cena: {text!r}")

    amount = Decimal(match.group(1).replace(" ", "").replace(",", "."))
    unit = normalise(match.group(2) or "").replace(".", "").replace(" ", "") or None
    return amount, unit


def retail_grosze(net: Decimal) -> int:
    gross = net * (Decimal(1) + MARGIN) * (Decimal(1) + VAT) * 100
    return int(gross.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def split_title(raw) -> tuple[str, str, bool]:
    """
    Rozbija komórkę z kolumny F na nazwę i opis.

    Dostawca oddziela je raz znakiem końca linii, a raz ciągiem spacji, więc podział idzie po
    obu naraz — i musi zdążyć przed `clean()`, które zwija spacje. Zwraca też informację, czy
    pozycja nie jest oznaczona jako niedostępna.

    Znaczniki („Nowość !", „Niedostępny !") zastępujemy ciągiem spacji, a nie jedną: w wierszach
    bez końca linii to one oddzielają nazwę od tego, co po nich następuje.
    """
    text = str(raw or "").replace(" ", " ")
    available = MARKER_UNAVAILABLE.search(text) is None
    text = MARKER_UNAVAILABLE.sub("   ", MARKER_NEW.sub("   ", text))

    lines = [line for line in text.split("\n") if line.strip()]
    if not lines:
        return "", "", available

    head, rest = lines[0].strip(), lines[1:]

    # Część wierszy nie ma końca linii — nazwa jest wtedy oddzielona od opisu ciągiem spacji.
    # Próg długości chroni nazwy z pojedynczą przerwą w środku („… 203 cm  perła").
    if len(head) > 60 and (split := re.search(r"\s{3,}", head)):
        head, tail = head[: split.start()], head[split.end():]
        rest.insert(0, tail)

    def tidy(value: str) -> str:
        return re.sub(r"\s+", " ", value).strip(" -–—")

    return tidy(head), "\n".join(filter(None, (tidy(line) for line in rest))), available


def build_description(
    name: str,
    long_text: str,
    specs: list[tuple[str, str]],
    unit: str | None,
) -> tuple[str, str]:
    """Zwraca (opis skrócony, opis pełny). Oba muszą być niepuste — wymaga tego domena."""
    body = long_text.strip()

    # Dostawca zapisuje opis jako dalszy ciąg zdania rozpoczętego nazwą („Pokrywa termiczna
    # 225 cm okrągła" + „ze wzmocnioną konstrukcją…"). Bez nazwy fragment czyta się jak urwany.
    if body[:1].islower():
        body = f"{name} {body}"

    spec_lines = [f"{label}: {value}" for label, value in specs]
    description = "\n\n".join(part for part in (body, "\n".join(spec_lines)) if part)

    # Gdy pozycja nie ma opisu, kafelek dostaje pierwszą daną techniczną — bez wiodącego
    # myślnika listy, którym dostawca zaczyna wyliczenia.
    lead = (body or (specs[0][1] if specs else "")).lstrip("-–—• ")
    sentence = re.split(r"(?<=[.!?])\s+", lead.replace("\n", " "), maxsplit=1)[0] if lead else ""
    short = sentence or lead

    if unit and (label := UNIT_LABELS.get(unit)):
        note = f"Cena za {label[1]}."
        short = f"{note} {short}".strip()
        description = f"{note}\n\n{description}".strip()

    return shorten(short, MAX_SHORT), shorten(description, MAX_DESCRIPTION)


def parse_sheet(path: Path) -> tuple[list[Product], list[str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook[SHEET].iter_rows(values_only=True))

    with zipfile.ZipFile(path) as archive:
        sheet_index = workbook.sheetnames.index(SHEET) + 1
        links = read_hyperlinks(archive, sheet_index)
        anchors = read_image_anchors(archive, sheet_index)

    products: list[Product] = []
    warnings: list[str] = []
    used_slugs: dict[str, int] = {}
    section: str | None = None
    labels: dict[int, str] = {}

    for index, row in enumerate(rows, start=1):
        cells = [clean(cell) for cell in row] + [""] * 16
        filled = [value for value in cells if value]

        if not filled:
            continue

        # Nagłówek tabeli — zapamiętuje etykiety kolumn z danymi technicznymi.
        if normalise(cells[COL_PRICE]).startswith("cena"):
            labels = {column: cells[column] for column in SPEC_COLS if cells[column]}
            continue

        # Wiersz z samą kolumną A to tytuł sekcji albo baner reklamowy.
        if len(filled) == 1 and cells[0]:
            title = normalise(cells[0])
            if title not in SECTION_BANNERS:
                section = title
                if title not in SECTION_CATEGORIES:
                    warnings.append(f"w. {index}: sekcja bez przypisanej kategorii — {cells[0]!r}")
            continue

        if not cells[COL_NAME]:
            continue

        category = ROW_CATEGORIES.get(index) or SECTION_CATEGORIES.get(section or "", "Accessories")

        try:
            net, unit = parse_price(row[COL_PRICE])
        except ValueError as error:
            warnings.append(f"w. {index}: {error} — pozycja pominięta")
            continue
        if net <= 0:
            warnings.append(f"w. {index}: cena zerowa — pozycja pominięta")
            continue

        name, body, available = split_title(row[COL_NAME])
        if not name:
            warnings.append(f"w. {index}: pusta nazwa — pozycja pominięta")
            continue

        prefix = SECTION_NAME_PREFIXES.get(section or "")
        if prefix and "wklad" not in normalise(name):
            name = f"{prefix} — {name}"
        if unit and (label := UNIT_LABELS.get(unit)):
            name = f"{name} — cena za {label[0]}"
        if not available:
            warnings.append(f"w. {index}: oznaczone jako niedostępne, import bez publikacji — {name!r}")

        specs = [
            (labels.get(column, "Dane techniczne"), cells[column])
            for column in SPEC_COLS
            if cells[column] and not cells[column].startswith("http")
        ]
        short_description, description = build_description(name, body, specs, unit)
        if not short_description:
            short_description = name

        slug = slugify(name) or f"produkt-{index}"
        if slug in used_slugs:
            used_slugs[slug] += 1
            slug = f"{slug}-{used_slugs[slug]}"[:MAX_SLUG]
        else:
            used_slugs[slug] = 1

        link = links.get(f"{chr(64 + COL_LINK + 1)}{index}")

        products.append(Product(
            slug=slug,
            name=shorten(name, MAX_NAME),
            short_description=short_description,
            description=description,
            price_grosze=retail_grosze(net),
            wholesale_net_grosze=int((net * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP)),
            category=category,
            sort_order=len(products) * 10 + 10,
            is_published=available,
            unit=unit,
            source_url=link,
            source_row=index,
            images=[target for _, target in anchors.get(index, [])][:MAX_IMAGES],
        ))

        if not anchors.get(index):
            warnings.append(f"w. {index}: brak zdjęcia — {name!r}")

    return products, warnings


def export_images(xlsx: Path, products: list[Product], media_dir: Path) -> dict[str, str]:
    """Konwertuje obrazy do WebP. Ten sam plik użyty w kilku wierszach zapisujemy raz."""
    media_dir.mkdir(parents=True, exist_ok=True)
    for stale in media_dir.glob("*.webp"):
        stale.unlink()

    wanted = {target for product in products for target in product.images}
    names: dict[str, str] = {}

    with zipfile.ZipFile(xlsx) as archive:
        for target in sorted(wanted):
            data = archive.read(target)
            digest = hashlib.sha1(data).hexdigest()[:16]
            name = f"katalog-{digest}.webp"
            names[target] = name

            destination = media_dir / name
            if destination.exists():
                continue

            with Image.open(io.BytesIO(data)) as image:
                if image.mode in ("RGBA", "LA", "P"):
                    background = Image.new("RGB", image.size, (255, 255, 255))
                    converted = image.convert("RGBA")
                    background.paste(converted, mask=converted.split()[-1])
                    image = background
                else:
                    image = image.convert("RGB")

                if image.width > IMAGE_MAX_WIDTH:
                    height = round(image.height * IMAGE_MAX_WIDTH / image.width)
                    image = image.resize((IMAGE_MAX_WIDTH, height), Image.LANCZOS)

                image.save(destination, "WEBP", quality=IMAGE_QUALITY, method=6)

    return names


def main() -> int:
    root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(description="Import hurtowego cennika do katalogu sklepu.")
    parser.add_argument("xlsx", type=Path, help="plik cennika (.xlsx)")
    parser.add_argument(
        "--out-json",
        type=Path,
        default=root / "backend/src/AkHouse.Infrastructure/Persistence/Data/shop-catalog.json",
    )
    parser.add_argument(
        "--out-media",
        type=Path,
        default=root / "backend/src/AkHouse.Infrastructure/Persistence/Data/CatalogMedia",
    )
    parser.add_argument("--skip-media", action="store_true", help="tylko JSON, bez konwersji zdjęć")
    args = parser.parse_args()

    products, warnings = parse_sheet(args.xlsx)

    names = {} if args.skip_media else export_images(args.xlsx, products, args.out_media)

    payload = {
        "source": args.xlsx.name,
        "marginPercent": float(MARGIN * 100),
        "vatPercent": float(VAT * 100),
        "products": [
            {
                "slug": product.slug,
                "name": product.name,
                "shortDescription": product.short_description,
                "description": product.description,
                "priceGrosze": product.price_grosze,
                "wholesaleNetGrosze": product.wholesale_net_grosze,
                "category": product.category,
                "sortOrder": product.sort_order,
                "isPublished": product.is_published,
                "unit": product.unit,
                "sourceUrl": product.source_url,
                "sourceRow": product.source_row,
                "images": [names[target] for target in product.images if target in names],
            }
            for product in products
        ],
    }

    args.out_json.parent.mkdir(parents=True, exist_ok=True)
    args.out_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    by_category: dict[str, int] = defaultdict(int)
    for product in products:
        by_category[product.category] += 1

    print(f"pozycji: {len(products)}   zdjęć: {len(names)}")
    for category, count in sorted(by_category.items(), key=lambda item: -item[1]):
        print(f"  {category:<16} {count}")
    if warnings:
        print(f"\nuwagi ({len(warnings)}):")
        for warning in warnings:
            print(f"  {warning}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
